"""文档解析模块。

支持 PDF / DOCX / TXT 三种格式，输出统一为纯文本。
"""
from pathlib import Path
from typing import Optional

from app.utils.logger import logger


def parse_document(file_path: str, file_type: Optional[str] = None) -> str:
    """解析文档，返回纯文本。

    Args:
        file_path: 文件路径
        file_type: 文件类型（pdf/docx/txt），None 时根据扩展名推断
    """
    path = Path(file_path)
    if file_type is None:
        file_type = path.suffix.lower().lstrip(".")

    file_type = file_type.lower()
    if file_type == "pdf":
        return _parse_pdf(path)
    elif file_type in ("docx", "doc"):
        return _parse_docx(path)
    elif file_type in ("txt", "md"):
        return _parse_txt(path)
    else:
        logger.warning(f"不支持的文件类型：{file_type}，尝试按文本读取")
        return _parse_txt(path)


def _parse_pdf(path: Path) -> str:
    """解析 PDF（普通 PDF，扫描件需 OCR，此处用 PyMuPDF 提取文本层）。"""
    try:
        import fitz  # pymupdf
    except ImportError:
        logger.warning("pymupdf 未安装，PDF 解析将返回空")
        return ""
    text_parts = []
    with fitz.open(path) as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts)


def _parse_docx(path: Path) -> str:
    """解析 DOCX。"""
    try:
        from docx import Document
    except ImportError:
        logger.warning("python-docx 未安装，DOCX 解析将返回空")
        return ""
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # 也提取表格内容
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return "\n".join(paragraphs)


def _parse_txt(path: Path) -> str:
    """解析 TXT，尝试多种编码。"""
    for encoding in ("utf-8", "gbk", "gb2312", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_text(errors="ignore")
